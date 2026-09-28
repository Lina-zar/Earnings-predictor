#!/usr/bin/env python3
"""
export_powerbi.py - build the data file for the Power BI dashboard.

Writes powerbi/earnings_dashboard.xlsx, one Excel table per sheet (a small star schema):
  fact_earnings  one row per labelled earnings report (5,380): outcome, surprise, stock
                 reaction, and the model's risk score (walk-forward, so 2019-2025 only)
  dim_company    one row per company: name, sector, index status, overall beat rate
  dim_date       one row per report date: year, quarter, month
  model_by_year  Brier score and skill vs the sector baseline, per model and year
  risk_deciles   actual not-beat rate and stock reaction by the model's risk decile

Percent columns end in _pct and are already in percent units (0.57 means 0.57%).

Inputs : data/processed/features_v2.csv, earnings_events.csv, reactions.csv,
         model_v1_predictions.csv, universe_companies.csv; company_names.csv
Run    : python export_powerbi.py      (needs openpyxl: pip install openpyxl)
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
PROC = ROOT / "data" / "processed"
OUT = ROOT / "powerbi" / "earnings_dashboard.xlsx"


def brier(y, p):
    return float(np.mean((np.asarray(p) - np.asarray(y)) ** 2))


def main():
    f = pd.read_csv(PROC / "features_v2.csv", parse_dates=["report_date"])
    ev = pd.read_csv(PROC / "earnings_events.csv", parse_dates=["report_date"])
    r = pd.read_csv(PROC / "reactions.csv", parse_dates=["report_date"])
    pr = pd.read_csv(PROC / "model_v1_predictions.csv", parse_dates=["report_date"])
    comp = pd.read_csv(PROC / "universe_companies.csv")
    names = pd.read_csv(ROOT / "company_names.csv")
    key = ["company_key", "report_date"]

    # ---- fact table
    fact = f[key + ["report_time", "eps_estimate", "outcome", "not_beat"]].merge(
        ev[key + ["eps_actual"]].drop_duplicates(key), on=key, how="left")
    fact = fact.merge(r[key + ["car_0_1", "car_0_2"]], on=key, how="left")
    fact = fact.merge(pr[key + ["logistic", "base_sector"]], on=key, how="left")
    fact["beat"] = 1 - fact["not_beat"]
    fact["surprise_eps"] = (fact["eps_actual"] - fact["eps_estimate"]).round(4)
    fact["surprise_pct"] = (100 * fact["surprise_eps"] / fact["eps_estimate"].abs()
                            .where(fact["eps_estimate"].abs() >= 0.05)).clip(-100, 100).round(2)
    fact["reaction_2d_pct"] = (100 * fact["car_0_1"]).round(3)
    fact["reaction_3d_pct"] = (100 * fact["car_0_2"]).round(3)
    fact["risk_score_pct"] = (100 * fact["logistic"]).round(2)
    fact["sector_baseline_pct"] = (100 * fact["base_sector"]).round(2)
    fact["year"] = fact["report_date"].dt.year
    has = fact["logistic"].notna()
    fact.loc[has, "risk_decile"] = fact[has].groupby("year")["logistic"].transform(
        lambda s: pd.qcut(s.rank(method="first"), 10, labels=False) + 1)
    fact.loc[has, "risk_band"] = fact.loc[has, "risk_decile"].map(
        lambda d: "1 Low (D1-D3)" if d <= 3 else ("3 High (D8-D10)" if d >= 8 else "2 Medium (D4-D7)"))
    fact["risk_decile"] = fact["risk_decile"].astype("Int64")
    fact["outcome"] = fact["outcome"].str.replace("in line", "In line").str.capitalize()
    fact = fact[["company_key", "report_date", "report_time", "outcome", "beat", "not_beat",
                 "eps_estimate", "eps_actual", "surprise_eps", "surprise_pct",
                 "reaction_2d_pct", "reaction_3d_pct", "risk_score_pct", "risk_decile",
                 "risk_band", "sector_baseline_pct"]].sort_values(key)

    # ---- company dimension
    stats = fact.groupby("company_key").agg(reports=("beat", "size"), beat_rate_pct=("beat", "mean"),
                                            avg_reaction_2d_pct=("reaction_2d_pct", "mean"))
    stats["beat_rate_pct"] = (100 * stats["beat_rate_pct"]).round(1)
    stats["avg_reaction_2d_pct"] = stats["avg_reaction_2d_pct"].round(2)
    dim_c = comp[["company_key", "sector", "status", "first_in", "last_out"]].merge(
        names, on="company_key", how="left").merge(stats, on="company_key", how="inner")
    dim_c["status"] = dim_c["status"].replace({"current member": "In the index today",
                                               "left after 2025": "Left after 2025",
                                               "left during window": "Left 2018-2025"})
    dim_c = dim_c[["company_key", "company_name", "sector", "status", "first_in", "last_out",
                   "reports", "beat_rate_pct", "avg_reaction_2d_pct"]]

    # ---- date dimension
    d = pd.DataFrame({"report_date": sorted(fact["report_date"].unique())})
    d["year"] = d["report_date"].dt.year
    d["quarter"] = "Q" + d["report_date"].dt.quarter.astype(str)
    d["year_quarter"] = d["year"].astype(str) + " " + d["quarter"]
    d["month"] = d["report_date"].dt.strftime("%b")
    d["month_number"] = d["report_date"].dt.month

    # ---- model performance by year
    rows = []
    names_m = {"base_overall": "Overall base rate", "base_sector": "Sector base rate",
               "logistic": "Logistic regression", "gbm": "Gradient-boosted trees"}
    for yr, g in pr.groupby("test_year"):
        ref = brier(g["not_beat"], g["base_sector"])
        for m, label in names_m.items():
            b = brier(g["not_beat"], g[m])
            rows.append({"year": yr, "model": label, "brier": round(b, 4),
                         "skill_vs_sector_pct": round(100 * (1 - b / ref), 2)})
    by_year = pd.DataFrame(rows)

    # ---- risk deciles
    dec = fact[fact["risk_decile"].notna()].groupby("risk_decile").agg(
        reports=("beat", "size"), predicted_risk_pct=("risk_score_pct", "mean"),
        actual_not_beat_pct=("not_beat", "mean"), avg_reaction_2d_pct=("reaction_2d_pct", "mean"))
    dec["actual_not_beat_pct"] *= 100
    dec = dec.round(2).reset_index()
    dec["risk_decile"] = dec["risk_decile"].astype(int)

    # ---- write workbook, each sheet as a named Excel table
    OUT.parent.mkdir(exist_ok=True)
    tables = {"fact_earnings": fact, "dim_company": dim_c, "dim_date": d,
              "model_by_year": by_year, "risk_deciles": dec}
    with pd.ExcelWriter(OUT, engine="openpyxl", date_format="YYYY-MM-DD",
                        datetime_format="YYYY-MM-DD") as xw:
        for name, t in tables.items():
            t.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(OUT)
    for name, t in tables.items():
        ws = wb[name]
        ref = f"A1:{get_column_letter(t.shape[1])}{t.shape[0] + 1}"
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
        for i, col in enumerate(t.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(col) + 2)
    wb.save(OUT)
    for name, t in tables.items():
        print(f"  {name:14} {len(t):>6,} rows")
    print(f"Saved {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
